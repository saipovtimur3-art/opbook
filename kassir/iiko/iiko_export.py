"""Разбор выгрузки OLAP-отчёта iikoOffice (Excel .xlsx или CSV).

Выгрузка «как есть» из iikoOffice: сверху шапка отчёта (название, период,
фильтры), потом строка заголовков, потом данные. В режиме сводной таблицы
iiko оставляет пустые ячейки у повторяющихся значений группировки и
вставляет строки «Итого» — всё это обрабатывается здесь.

На выходе список строк с ключами:
    day, department, cashier, session, check, group, dish, amount, revenue
"""
import csv
import datetime as dt
import io
import re

# канонический ключ -> варианты названий колонки в iikoOffice (нижний регистр)
ALIASES = {
    "day":        ["учетный день", "учётный день", "дата открытия", "дата"],
    "department": ["торговое предприятие", "подразделение", "ресторан", "точка"],
    "cashier":    ["кассир", "сотрудник", "официант заказа", "официант"],
    "session":    ["номер кассовой смены", "кассовая смена", "номер смены"],
    "check":      ["номер чека", "номер заказа", "чек", "заказ"],
    "group":      ["группа блюда", "группа товара", "группа", "категория блюда", "категория"],
    "dish":       ["блюдо", "товар", "наименование"],
    "amount":     ["количество блюд", "количество товаров", "количество", "кол-во"],
    "revenue":    ["сумма со скидкой", "сумма со скидками", "сумма после скидки", "выручка", "сумма"],
}
REQUIRED = ["cashier", "check", "dish", "amount", "revenue"]
FILL_DOWN = ["day", "department", "cashier", "session", "check", "group"]
TOTAL_RE = re.compile(r"^\s*(итого|всего|total)", re.I)


class ExportError(Exception):
    pass


def _norm(v):
    return re.sub(r"\s+", " ", str(v or "")).strip().lower().replace("ё", "е")


def _match_headers(cells):
    """{канонический ключ: индекс колонки}. Точное совпадение важнее префикса,
    префикс важнее вхождения; каждая колонка используется один раз."""
    heads = [_norm(c) for c in cells]
    used, out = set(), {}
    for key, aliases in ALIASES.items():
        best = None
        for prio, test in enumerate((lambda h, a: h == a, lambda h, a: h.startswith(a), lambda h, a: a in h)):
            for a in aliases:
                for i, h in enumerate(heads):
                    if i not in used and h and test(h, _norm(a)):
                        best = i
                        break
                if best is not None:
                    break
            if best is not None:
                break
        if best is not None:
            out[key] = best
            used.add(best)
    return out


def _num(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace("\xa0", "").replace(" ", "").replace("₽", "").replace("р.", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def _day(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    s = str(v or "").strip()
    for fmt in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return dt.datetime.strptime(s.split(" ")[0], fmt).date().isoformat()
        except ValueError:
            pass
    return s


def _read_table(data, filename):
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    if name.endswith(".csv") or name.endswith(".txt"):
        text = data.decode("utf-8-sig", errors="replace")
        if text.count("�") > 5:
            text = data.decode("cp1251", errors="replace")
        dialect = csv.Sniffer().sniff(text[:4000], delimiters=";,\t")
        return [r for r in csv.reader(io.StringIO(text), dialect)]
    if name.endswith(".xls"):
        raise ExportError("Старый формат .xls не поддерживается. В iikoOffice выбери «Экспорт в Excel (.xlsx)» или сохрани файл как .xlsx.")
    raise ExportError("Нужен файл .xlsx или .csv — выгрузка OLAP-отчёта из iikoOffice.")


def parse_export(data, filename):
    table = _read_table(data, filename)

    # строка заголовков — та, где узнаётся больше всего колонок
    best_i, best_map = None, {}
    for i, row in enumerate(table[:40]):
        m = _match_headers(row)
        if len(m) > len(best_map):
            best_i, best_map = i, m
    missing = [k for k in REQUIRED if k not in best_map]
    if best_i is None or missing:
        names = {"cashier": "Кассир", "check": "Номер чека", "dish": "Блюдо", "amount": "Количество блюд", "revenue": "Сумма со скидкой"}
        raise ExportError("В отчёте не найдены колонки: " + ", ".join(names[k] for k in missing) +
                          ". Проверь, что отчёт построен по инструкции.")

    rows, prev = [], {}
    for raw in table[best_i + 1:]:
        if not raw or all(c is None or str(c).strip() == "" for c in raw):
            continue
        get = lambda k: raw[best_map[k]] if k in best_map and best_map[k] < len(raw) else None
        if any(TOTAL_RE.match(str(c or "")) for c in raw):
            continue
        r = {}
        for k in ALIASES:
            v = get(k)
            if k in FILL_DOWN and (v is None or str(v).strip() == ""):
                v = prev.get(k)
            r[k] = v
        if not str(r.get("dish") or "").strip():      # промежуточный итог группы
            continue
        prev = r
        rows.append({
            "day": _day(r["day"]),
            "department": str(r["department"] or "").strip(),
            "cashier": str(r["cashier"] or "").strip() or "Без кассира",
            "session": str(r["session"] or "").strip(),
            "check": str(r["check"] or "").strip(),
            "group": str(r["group"] or "").strip(),
            "dish": str(r["dish"]).strip(),
            "amount": _num(r["amount"]),
            "revenue": _num(r["revenue"]),
        })
    if not rows:
        raise ExportError("В отчёте нет строк с продажами.")
    return rows
